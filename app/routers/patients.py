import io
import os
import io
import logging
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session

from pydantic import BaseModel

from app.auth import get_current_user, require_role
from app.audit_log import log_action
from app.db.database import get_db
from app.db.models import Invoice, Patient, PatientDocument, User, UserRole
from app.email_service import send_email_with_attachment
from app.storage import abs_path, get_patient_docs_path, unique_name


class DocEmailRequest(BaseModel):
    email: str | None = None


def _doc_abs_path(db: Session, doc: PatientDocument) -> str | None:
    """Ruta absoluta de un documento de paciente a partir del nombre relativo."""
    if not doc.filepath:
        return None
    if os.path.isabs(doc.filepath):
        return doc.filepath  # compatibilidad con datos antiguos
    return abs_path(get_patient_docs_path(db, ensure=False, check_mount=False), doc.filepath)

from dotenv import load_dotenv

load_dotenv()

router = APIRouter(prefix="/api/patients", tags=["patients"])


@router.get("")
def list_patients(
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    search: str = Query("", alias="q"),
    pending: bool = Query(False),
    show_inactive: bool = Query(False),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Patient)
    if not show_inactive:
        query = query.filter(Patient.is_active == True)
    else:
        query = query.filter(Patient.is_active == False)
    if pending:
        query = query.join(Invoice, Invoice.patient_id == Patient.id).filter(Invoice.is_paid == False).distinct()
    if search:
        term = f"%{search}%"
        query = query.filter(
            or_(
                Patient.first_name.ilike(term),
                Patient.last_name.ilike(term),
                Patient.phone.ilike(term),
            )
        )
    total = query.count()
    patients = query.order_by(Patient.last_name, Patient.first_name).offset((page - 1) * size).limit(size).all()
    return {
        "patients": [
            {
                "id": p.id,
                "first_name": p.first_name,
                "last_name": p.last_name,
                "phone": p.phone,
                "email": p.email,
                "allergies": p.allergies,
                "is_active": p.is_active,
            }
            for p in patients
        ],
        "total": total,
        "page": page,
        "pages": (total + size - 1) // size,
    }


@router.post("", status_code=201)
def create_patient(
    first_name: str = Form(...),
    last_name: str = Form(...),
    phone: str = Form(...),
    email: str = Form(None),
    address: str = Form(None),
    birth_date: str = Form(None),
    dni: str = Form(None),
    allergies: str = Form(None),
    notes: str = Form(None),
    motivo_consulta: str = Form(None),
    anamnesis: str = Form(None),
    tratamiento_contraindicaciones: str = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    patient = Patient(
        first_name=first_name.strip(),
        last_name=last_name.strip(),
        phone=phone.strip(),
        email=email.strip() if email else None,
        address=address.strip() if address else None,
        birth_date=date.fromisoformat(birth_date) if birth_date else None,
        dni=dni.strip() if dni else None,
        allergies=allergies.strip() if allergies else None,
        notes=notes.strip() if notes else None,
        motivo_consulta=motivo_consulta.strip() if motivo_consulta else None,
        anamnesis=anamnesis.strip() if anamnesis else None,
        tratamiento_contraindicaciones=tratamiento_contraindicaciones.strip() if tratamiento_contraindicaciones else None,
    )
    db.add(patient)
    db.commit()
    db.refresh(patient)
    log_action(db, current_user.id, "CREAR", "PACIENTE", patient.id, f"{patient.first_name} {patient.last_name}")
    return {"id": patient.id, "message": "Paciente creado"}


@router.get("/{patient_id}")
def get_patient(
    patient_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    if not patient:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")
    return {
        "id": patient.id,
        "first_name": patient.first_name,
        "last_name": patient.last_name,
        "phone": patient.phone,
        "email": patient.email,
        "address": patient.address,
        "birth_date": patient.birth_date.isoformat() if patient.birth_date else None,
        "dni": patient.dni,
        "allergies": patient.allergies,
        "notes": patient.notes,
        "motivo_consulta": patient.motivo_consulta,
        "anamnesis": patient.anamnesis,
        "tratamiento_contraindicaciones": patient.tratamiento_contraindicaciones,
        "created_at": patient.created_at.isoformat() if patient.created_at else None,
        "documents": [
            {"id": d.id, "filename": d.filename, "description": d.description, "uploaded_at": d.uploaded_at.isoformat() if d.uploaded_at else None}
            for d in patient.documents
        ],
    }


@router.put("/{patient_id}")
def update_patient(
    patient_id: int,
    first_name: str = Form(...),
    last_name: str = Form(...),
    phone: str = Form(...),
    email: str = Form(None),
    address: str = Form(None),
    birth_date: str = Form(None),
    dni: str = Form(None),
    allergies: str = Form(None),
    notes: str = Form(None),
    motivo_consulta: str = Form(None),
    anamnesis: str = Form(None),
    tratamiento_contraindicaciones: str = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    if not patient:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")

    patient.first_name = first_name.strip()
    patient.last_name = last_name.strip()
    patient.phone = phone.strip()
    patient.email = email.strip() if email else None
    patient.address = address.strip() if address else None
    patient.birth_date = date.fromisoformat(birth_date) if birth_date else None
    patient.dni = dni.strip() if dni else None
    patient.allergies = allergies.strip() if allergies else None
    patient.notes = notes.strip() if notes else None
    patient.motivo_consulta = motivo_consulta.strip() if motivo_consulta else None
    patient.anamnesis = anamnesis.strip() if anamnesis else None
    patient.tratamiento_contraindicaciones = tratamiento_contraindicaciones.strip() if tratamiento_contraindicaciones else None

    db.commit()
    log_action(db, current_user.id, "EDITAR", "PACIENTE", patient.id, f"{patient.first_name} {patient.last_name}")
    return {"message": "Paciente actualizado"}


@router.get("/{patient_id}/documents")
def list_documents(
    patient_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Lista TODOS los documentos del paciente (subidos, consentimientos firmados,
    revocaciones... todo unificado), ordenados por fecha descendente."""
    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    if not patient:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")
    docs = (
        db.query(PatientDocument)
        .filter(PatientDocument.patient_id == patient_id)
        .order_by(PatientDocument.uploaded_at.desc())
        .all()
    )
    return [
        {
            "id": d.id,
            "filename": d.filename,
            "description": d.description,
            "uploaded_at": d.uploaded_at.isoformat() if d.uploaded_at else None,
            # is_pdf permite al front decidir: abrir en pestaña (PDF) o descargar (otro).
            "is_pdf": (d.filename or "").lower().endswith(".pdf"),
        }
        for d in docs
    ]


@router.post("/{patient_id}/documents", status_code=201)
def upload_document(
    patient_id: int,
    files: list[UploadFile] = File(...),
    description: str = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Sube uno o varios documentos a la vez. Cada fichero se guarda con un nombre
    único en disco; en BD se conserva el nombre original para mostrarlo."""
    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    if not patient:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")

    base = get_patient_docs_path(db)  # exige USB montado (lanza 503 si no)
    saved = 0
    for file in files:
        ext = os.path.splitext(file.filename)[1] if file.filename else ""
        stored_name = unique_name(f"p{patient_id}", ext)
        filepath = os.path.join(base, stored_name)
        with open(filepath, "wb") as f:
            f.write(file.file.read())
        doc = PatientDocument(
            patient_id=patient_id,
            filename=file.filename or stored_name,
            filepath=stored_name,  # ruta relativa (base en Config)
            description=description.strip() if description else None,
        )
        db.add(doc)
        saved += 1
    db.commit()
    return {"message": f"{saved} documento(s) subido(s)", "count": saved}


@router.get("/{patient_id}/documents/{doc_id}/download")
def download_document(
    patient_id: int,
    doc_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    doc = db.query(PatientDocument).filter(
        PatientDocument.id == doc_id, PatientDocument.patient_id == patient_id
    ).first()
    path = _doc_abs_path(db, doc) if doc else None
    if not doc or not path or not os.path.exists(path):
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    return FileResponse(path, filename=doc.filename)


@router.get("/{patient_id}/documents/{doc_id}/view")
def view_document(
    patient_id: int,
    doc_id: int,
    token: str = Query(""),
    db: Session = Depends(get_db),
):
    from app.auth import verify_token
    if not token:
        raise HTTPException(status_code=401, detail="Token requerido")
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="Token invalido")
    doc = db.query(PatientDocument).filter(
        PatientDocument.id == doc_id, PatientDocument.patient_id == patient_id
    ).first()
    path = _doc_abs_path(db, doc) if doc else None
    if not doc or not path or not os.path.exists(path):
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    import mimetypes
    media_type = mimetypes.guess_type(doc.filename)[0] or "application/octet-stream"
    return FileResponse(path, media_type=media_type, filename=doc.filename, content_disposition_type="inline")


@router.delete("/{patient_id}/documents/{doc_id}", status_code=200)
def delete_document(
    patient_id: int,
    doc_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    doc = db.query(PatientDocument).filter(
        PatientDocument.id == doc_id, PatientDocument.patient_id == patient_id
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    path = _doc_abs_path(db, doc)
    if path and os.path.exists(path):
        os.remove(path)
    db.delete(doc)
    db.commit()
    log_action(db, current_user.id, "ELIMINAR", "DOCUMENTO", doc_id, f"Paciente {patient_id}")
    return {"message": "Documento eliminado"}


@router.post("/{patient_id}/documents/{doc_id}/email")
def email_document(
    patient_id: int,
    doc_id: int,
    body: DocEmailRequest = DocEmailRequest(),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Envía un documento del paciente por email (cualquier tipo de fichero)."""
    doc = db.query(PatientDocument).filter(
        PatientDocument.id == doc_id, PatientDocument.patient_id == patient_id
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado")

    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    to_email = body.email or (patient.email if patient else None)
    if not to_email:
        raise HTTPException(status_code=400, detail="El paciente no tiene email registrado")

    path = _doc_abs_path(db, doc)
    if not path or not os.path.exists(path):
        raise HTTPException(status_code=404, detail="Fichero no encontrado")

    clinic = os.getenv("CLINIC_NAME", "Santé")
    try:
        send_email_with_attachment(
            to_email=to_email,
            subject=f"Documento - {clinic}",
            body=f"Adjunto encontrará el documento solicitado.\n\nUn saludo,\n{clinic}",
            attachment_path=path,
            attachment_filename=doc.filename,
        )
    except Exception as e:
        logging.error(f"Error enviando email de documento: {e}")
        raise HTTPException(status_code=500, detail=f"Error al enviar email: {str(e)}")

    return {"message": f"Email enviado a {to_email}"}


@router.patch("/{patient_id}/deactivate")
def toggle_patient_active(
    patient_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(UserRole.ADMIN, UserRole.RECEPTION)),
):
    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    if not patient:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")
    patient.is_active = not patient.is_active
    db.commit()
    action = "REACTIVAR" if patient.is_active else "DESACTIVAR"
    log_action(db, current_user.id, action, "PACIENTE", patient.id, f"{patient.first_name} {patient.last_name}")
    return {"message": f"Paciente {'activado' if patient.is_active else 'desactivado'}", "is_active": patient.is_active}


@router.post("/import-excel", status_code=200)
def import_patients_from_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(UserRole.ADMIN, UserRole.RECEPTION)),
):
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no está instalado")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file.file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo Excel")

    if "CLIENTES" not in wb.sheetnames:
        raise HTTPException(status_code=400, detail="El Excel no contiene la hoja CLIENTES")

    ws = wb["CLIENTES"]
    headers = [str(ws.cell(1, c).value or "").strip().upper() for c in range(1, ws.max_column + 1)]

    def col(name):
        """Devuelve el índice (1-based) de una columna por nombre parcial."""
        for i, h in enumerate(headers):
            if name in h:
                return i + 1
        return None

    c_paciente = col("PACIENTE")
    c_nif      = col("NIF")
    c_dir      = col("DIRECCI")
    c_cp       = col("POSTAL")
    c_ciudad   = col("POBLACI")
    c_prov     = col("PROVINCIA")
    c_tel      = col("TEL")
    c_mail     = col("MAIL")

    if not c_paciente:
        raise HTTPException(status_code=400, detail="No se encontró la columna PACIENTE")

    created, skipped, errors, to_review = 0, 0, [], []

    for row in range(2, ws.max_row + 1):
        raw_name = ws.cell(row, c_paciente).value
        if not raw_name or str(raw_name).startswith("="):
            continue
        raw_name = str(raw_name).strip()
        if not raw_name:
            continue

        # Parsear "Apellidos, Nombre" — si no hay coma se guarda tal cual y se marca para revisar
        needs_review = False
        if "," in raw_name:
            last_name, _, first_name = raw_name.partition(",")
            last_name  = last_name.strip()
            first_name = first_name.strip()
        else:
            first_name = raw_name
            last_name  = ""
            needs_review = True
            to_review.append(f"Fila {row} ({raw_name}): nombre sin coma separadora, revisar nombre/apellidos")

        if not first_name:
            errors.append(f"Fila {row}: nombre vacío")
            continue

        phone = str(ws.cell(row, c_tel).value or "").strip() if c_tel else ""
        dni   = str(ws.cell(row, c_nif).value or "").strip() if c_nif else None
        email = str(ws.cell(row, c_mail).value or "").strip() if c_mail else None
        addr  = str(ws.cell(row, c_dir).value or "").strip() if c_dir else None
        cp    = str(ws.cell(row, c_cp).value or "").strip() if c_cp else None
        city  = str(ws.cell(row, c_ciudad).value or "").strip() if c_ciudad else None
        prov  = str(ws.cell(row, c_prov).value or "").strip() if c_prov else None

        # Sin teléfono: asignar 000000000 y marcar para revisar
        if not phone:
            phone = "000000000"
            to_review.append(f"Fila {row} ({raw_name}): sin teléfono, asignado 000000000")

        # Construir dirección completa con todos los campos
        address_parts = [p for p in [addr, cp, city, prov] if p]
        full_address = ", ".join(address_parts) or None

        # Comprobar duplicado por teléfono o DNI (solo si el teléfono no es el placeholder)
        dup_query = db.query(Patient).filter(Patient.phone == phone) if phone != "000000000" else None
        if dni:
            dup_query = db.query(Patient).filter(
                or_(Patient.phone == phone, Patient.dni == dni)
            ) if phone != "000000000" else db.query(Patient).filter(Patient.dni == dni)
        if dup_query and dup_query.first():
            skipped += 1
            continue

        patient = Patient(
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            email=email or None,
            address=full_address,
            dni=dni or None,
        )
        db.add(patient)
        created += 1

    db.commit()
    log_action(db, current_user.id, "IMPORTAR", "PACIENTE", None, f"{created} creados, {skipped} omitidos")
    return {"created": created, "skipped": skipped, "errors": errors, "to_review": to_review}
def toggle_patient_active(
    patient_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(UserRole.ADMIN, UserRole.RECEPTION)),
):
    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    if not patient:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")
    patient.is_active = not patient.is_active
    db.commit()
    action = "REACTIVAR" if patient.is_active else "DESACTIVAR"
    log_action(db, current_user.id, action, "PACIENTE", patient.id, f"{patient.first_name} {patient.last_name}")
    return {"message": f"Paciente {'activado' if patient.is_active else 'desactivado'}", "is_active": patient.is_active}
